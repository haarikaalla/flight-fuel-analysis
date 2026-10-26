"""Build the Springer Nature style review manuscript from verified outputs."""

from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
HERE = Path(__file__).resolve().parent
OUTPUT = HERE / "PI-MCS_Springer_Nature_manuscript.docx"
DATA = json.loads((ROOT / "research_q1/outputs/metadata-baselines.json").read_text())
INTERPRET = json.loads((ROOT / "research_q1/outputs/metadata-interpretability.json").read_text())
LABELS = {
    "hist_gradient_boosting": "Histogram gradient boosting",
    "histgb_ensemble": "Histogram boosting ensemble",
    "mlp": "Multilayer perceptron",
    "aircraft_type_median_rate": "Aircraft type median rate",
    "quantile_regression": "Quantile regression",
    "lasso": "Lasso",
    "ridge": "Ridge",
}


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    element = OxmlElement("w:shd")
    element.set(qn("w:fill"), fill)
    tc_pr.append(element)


def border(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        item = OxmlElement(f"w:{edge}")
        item.set(qn("w:val"), "single")
        item.set(qn("w:sz"), "4")
        item.set(qn("w:color"), "D9D9D9")
        borders.append(item)
    tc_pr.append(borders)


def table(doc, headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for j, width in enumerate(widths):
        t.columns[j].width = Inches(width)
    for j, head in enumerate(headers):
        cell = t.rows[0].cells[j]
        cell.text = head
        shade(cell, "26364A")
        for run in cell.paragraphs[0].runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.size = Pt(8)
    tr_pr = t.rows[0]._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    tr_pr.append(repeat)
    for i, values in enumerate(rows):
        cells = t.add_row().cells
        for j, value in enumerate(values):
            cells[j].text = str(value)
            if i % 2:
                shade(cells[j], "F1F4F7")
    for row in t.rows:
        for j, cell in enumerate(row.cells):
            cell.width = Inches(widths[j])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            border(cell)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.space_before = Pt(0)
                if j:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in paragraph.runs:
                    run.font.size = Pt(8)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def caption(doc, label, text):
    p = doc.add_paragraph(style="Caption")
    p.add_run(f"{label}. ").bold = True
    p.add_run(text)


def equation(doc, name, width=5.8):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    p.add_run().add_picture(str(HERE / "generated/equations" / f"{name}.png"), width=Inches(width))


def figure(doc, filename, label, description, width=5.7):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(HERE / "generated/figures" / filename), width=Inches(width))
    caption(doc, label, description)


def paragraph(doc, text):
    return doc.add_paragraph(text)


def heading(doc, text, level=1):
    return doc.add_heading(text, level=level)


def build():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(.8)
    section.bottom_margin = Inches(.8)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    styles = doc.styles
    for name in ("Normal", "Title", "Heading 1", "Heading 2", "Caption"):
        style = styles[name]
        style.font.name = "Times New Roman"
        style.font.color.rgb = RGBColor(0, 0, 0)
    styles["Normal"].font.size = Pt(10.5)
    styles["Normal"].paragraph_format.space_after = Pt(6)
    styles["Normal"].paragraph_format.line_spacing = 1.12
    styles["Title"].font.size = Pt(17)
    styles["Title"].font.bold = True
    styles["Title"].paragraph_format.space_after = Pt(10)
    title_properties = styles["Title"]._element.get_or_add_pPr()
    for title_border in title_properties.findall(qn("w:pBdr")):
        title_properties.remove(title_border)
    styles["Heading 1"].font.size = Pt(12)
    styles["Heading 1"].font.bold = True
    styles["Heading 1"].paragraph_format.space_before = Pt(12)
    styles["Heading 1"].paragraph_format.space_after = Pt(4)
    styles["Heading 2"].font.size = Pt(10.5)
    styles["Heading 2"].font.bold = True
    styles["Heading 2"].paragraph_format.space_before = Pt(9)
    styles["Caption"].font.size = Pt(8.5)
    styles["Caption"].font.italic = True

    doc.add_paragraph("Physics Informed Mondrian Conformal Selection for Aircraft Fuel Estimates", style="Title")
    paragraph(doc, "A reproducible development study using the EUROCONTROL PRC 2025 fuel benchmark")
    paragraph(doc, "Anonymized review manuscript | 26 October 2026")
    heading(doc, "Abstract")
    paragraph(doc, "Aircraft fuel estimates from open trajectories support emissions accounting and operational analysis, but sparse fuel labels, incomplete trajectory context, and unequal aircraft representation make prediction reliability difficult to assess. We formulate physics informed Mondrian conformal selection (PI-MCS), which joins a training-only physics residual predictor, aircraft by phase uncertainty calibration, and a mixed integer risk budget for deciding which estimates to release. Our methodological contribution is the joint decision formulation and its planned evaluation across temporal, route, and aircraft shifts; its individual building blocks are established methods. On a verified metadata-only development benchmark of 193,275 real fuel intervals from 15,761 flights, histogram gradient boosting attained 79.22 kg mean absolute error on 37,456 October intervals (95% route-cluster bootstrap interval 67.15 to 95.25 kg), versus 185.23 kg for ridge regression. A 90% split conformal interval around boosting predictions covered 90.92% of October labels with 293.20 kg mean width. Refit ablations show interval duration is the dominant metadata feature. These empirical results do not test PI-MCS itself; October labels had been inspected in the inherited project. We state the required trajectory, weather, multi-year, and independent confirmation analyses to turn this development study into a confirmatory evaluation.")
    paragraph(doc, "Keywords: aircraft fuel burn; prediction intervals; conformal prediction; physics informed learning; selective prediction; mixed integer optimization")

    heading(doc, "1 Introduction")
    paragraph(doc, "Fuel burn estimates affect emissions inventories, trajectory assessment, and comparisons of operational efficiency. The 2025 PRC challenge made ACARS fuel-on-board differences available alongside OpenSky trajectories, enabling public evaluation while retaining noisy labels and geographic and fleet selection effects [1]. Open data systems must therefore decide both how much fuel was likely consumed and when an estimate is reliable enough to release.")
    paragraph(doc, "Recent aviation studies combine OpenAP, weather, trajectory and machine learning [2, 3]. General conformal methods estimate interval uncertainty [4-7], while selective prediction and robust optimization address downstream decisions [8, 9]. These components have not yet been jointly demonstrated on this benchmark with aircraft by phase calibration, explicit risk budgets, representation audits, and grouped shift validation. That is the targeted research gap; no claim of a new physics law, conformal theorem, or optimizer is made.")
    paragraph(doc, "This study contributes a pre-specified decision formulation and executable validation framework; seven leakage-aware real-data metadata baselines with route-cluster uncertainty; and an explicit evidence boundary that separates observed performance from proposed full-system experiments. The practical question is whether uncertainty-aware selection can reduce large fuel-estimate errors without systematically excluding long intervals or minority aircraft classes. The current evidence establishes a development baseline, while that causal comparison remains to be measured.")

    heading(doc, "2 Data and study design")
    paragraph(doc, "The PRC 2025 release contains 15,761 flights, 193,275 labeled intervals, and 27 aircraft types between April and October 2025 [1]. Fuel labels are differences between irregular ACARS reports, with unit ambiguity resolved by broad OpenAP plausibility filtering in the source dataset. Airline identities are anonymized. We checksum-verified the six published flight-list and fuel-label files used here and kept raw records outside version control.")
    paragraph(doc, "For the reported metadata experiment, 11,037 flights and 131,530 intervals from April to August form training, 1,888 flights and 24,289 September intervals form calibration, and 2,836 flights and 37,456 October intervals form evaluation. All intervals of a flight stay together. The reported predictors use interval duration, whole-flight timing, time of day, and aircraft type; no trajectory vectors, weather, or phase labels enter these scores. October labels had already been examined in the inherited project; the evaluation is exploratory, not a locked external test.")
    caption(doc, "Table 1", "Verified cohort sizes for the exploratory metadata experiment.")
    table(doc, ["Period", "Flights", "Intervals", "Role"],
          [["April to August 2025", "11,037", "131,530", "Model fitting"],
           ["September 2025", "1,888", "24,289", "Interval calibration"],
           ["October 2025", "2,836", "37,456", "Exploratory evaluation"]], [2.1, 1.0, 1.1, 2.2])

    heading(doc, "3 Proposed PI MCS method")
    heading(doc, "3.1 Physics residual predictor", 2)
    paragraph(doc, "For each segment, a physical fuel prior combines ISA density, wind-corrected true airspeed, drag, cruise Breguet burn, and climb or descent energy balance. Takeoff mass, thrust specific fuel consumption, zero-lift drag and Oswald efficiency are fit with bounded training-fold least squares by aircraft type. A residual learner then corrects the physical estimate using only inputs available at prediction time:")
    equation(doc, "residual", 2.5)
    paragraph(doc, "In cruise, the prior includes the Breguet relation below, where m₀ is initial mass, cT is thrust specific fuel consumption, R is distance, V is true airspeed, and L/D is lift-to-drag ratio. Missing mass, wind or geometry invokes declared fallback parameters rather than test-label fitting.")
    equation(doc, "breguet", 3.8)
    heading(doc, "3.2 Aircraft by phase uncertainty calibration", 2)
    paragraph(doc, "A separate out-of-flight-fold model estimates error scale. On a disjoint calibration period, normalized absolute errors define nonconformity scores. Aircraft type by flight phase cells use a finite-sample quantile when at least 50 calibration examples exist; otherwise calibration falls back to aircraft type, phase, then global. Sparse fallback counts must be reported alongside subgroup coverage. Physical clipping preserves the feasible fuel range when tank capacity is known.")
    equation(doc, "score", 2.8)
    equation(doc, "interval", 4.5)
    paragraph(doc, "The quantile rank is ceil((n+1)(1−α)); if this exceeds n, the bound is infinite before physical clipping. Ordinary split conformal provides marginal coverage under exchangeability. Hierarchical fallback does not establish universal conditional coverage, and estimated density-ratio weighting has no oracle guarantee [4-7].")
    heading(doc, "3.3 Risk budget and released estimates", 2)
    paragraph(doc, "For N candidate intervals, a calibrated error-event probability rᵢ and operational utility uᵢ define a selection decision zᵢ. The optimization maximizes released utility subject to total expected error-event budget B and minimum coverage fraction m. Probabilities are calibrated on data separate from model fitting. The risk budget controls a model-based expectation; it is not an empirical guarantee of selected-set coverage.")
    equation(doc, "selection", 5.1)
    paragraph(doc, "We also implement a Bertsimas–Sim uncertainty set with deviation dᵢ and robustness budget Γ. This creates an auditable utility versus violation trade-off under perturbed risk estimates [9].")
    equation(doc, "robust", 4.8)
    paragraph(doc, "Exact mixed integer solutions are compared at matched budgets with risk ranking, utility-to-risk greedy selection, Lagrangian relaxation, and conformal selection when its dependence assumptions hold. Solver optimality and statistical calibration are reported separately.")

    heading(doc, "4 Validation and analysis plan")
    paragraph(doc, "The primary future tests are forward temporal, unseen route, unseen aircraft type, and unseen aircraft family partitions; grouped inner folds and flight-level disjointness prevent leakage. Fixed seeds, source hashes, a declared calibration period, and frozen analysis settings are included in the repository. The full experiment will compare the physical prior, residual learners, ridge, lasso, boosted trees, MLP, tree ensemble, and quantile prediction under an equal logged tuning budget.")
    paragraph(doc, "Point outcomes are MAE, RMSE and bias; interval outcomes are marginal and aircraft by phase coverage, width, and Winkler score. Decision outcomes are released fraction, utility, risk-budget violation, subgroup representation, solver time, and optimality gap. Route or aircraft-cluster bootstrap intervals, paired effect sizes, and Holm-adjusted multiple comparisons are pre-specified. Ablations remove physics, weather, trajectory, quality and calibration components. Missingness, wind, mass, sensor coverage and label-noise sensitivity are assessed separately. No unavailable result is filled with simulated evidence.")

    heading(doc, "5 Verified exploratory results")
    ordered = sorted(DATA["models"], key=lambda name: DATA["models"][name]["regression"]["mae_kg"])
    rows = []
    for key in ordered:
        entry = DATA["models"][key]
        metric = entry["regression"]
        ci = entry["mae_route_cluster_ci"]["cluster_bootstrap_95pct_ci"]
        rows.append([LABELS[key], f"{metric['mae_kg']:.2f}", f"{ci[0]:.2f}–{ci[1]:.2f}",
                     f"{metric['rmse_kg']:.2f}", f"{metric['r2']:.3f}"])
    caption(doc, "Table 2", "October 2025 point prediction. Confidence intervals resample 471 complete route clusters (2,000 replicates).")
    table(doc, ["Method", "MAE kg", "95% CI kg", "RMSE kg", "R²"], rows, [2.25, .8, 1.45, 1.0, .8])
    paragraph(doc, "Histogram gradient boosting achieved 79.22 kg MAE (95% route-cluster interval 67.15 to 95.25 kg). Ridge regression reached 185.23 kg. The paired MAE difference between the boosting ensemble and a single boosted model was 0.26 kg (95% cluster interval −0.17 to 0.74 kg), which does not support a clear ordering. This ensemble comprises boosted trees, not deep neural networks.")
    figure(doc, "metadata_baselines.png", "Figure 1", "Exploratory October MAE with route-cluster confidence intervals. The shared labels and overlapping intervals limit claims of model superiority.")

    interval_rows = []
    for key in ordered:
        metric = DATA["models"][key]["split_conformal_90"]
        interval_rows.append([LABELS[key], f"{100*metric['coverage']:.2f}",
                              f"{metric['mean_width_kg']:.2f}", f"{metric['winkler_score']:.2f}"])
    cqr = DATA["models"]["quantile_regression"]["cqr_90"]
    interval_rows.append(["Conformalized quantile regression", f"{100*cqr['coverage']:.2f}",
                          f"{cqr['mean_width_kg']:.2f}", f"{cqr['winkler_score']:.2f}"])
    caption(doc, "Table 3", "Nominal 90% September-calibrated intervals evaluated on October labels.")
    table(doc, ["Method", "Coverage %", "Mean width kg", "Winkler score"],
          interval_rows, [2.5, 1.1, 1.4, 1.4])
    paragraph(doc, "Boosting plus split conformal obtained 90.92% empirical coverage with 293.20 kg mean width. CQR covered 89.68% with 304.85 kg mean width and a lower Winkler score (548.03 versus 858.77); these metrics reward different aspects of interval quality and should be read together.")
    figure(doc, "metadata_coverage_width.png", "Figure 2", "Coverage versus width for the seven split conformal metadata baselines. The dashed line marks nominal 90% coverage.")

    ablations = [row for row in INTERPRET["refit_feature_group_ablation"] if row["ablation"] != "none"]
    caption(doc, "Table 4", "Refit feature-group ablation relative to the full histogram boosting model (79.22 kg MAE).")
    table(doc, ["Removed feature group", "MAE kg", "Change kg"],
          [[row["ablation"].replace("_", " ").capitalize(), f"{row['mae_kg']:.2f}",
            f"{row['delta_mae_kg']:+.2f}"] for row in ablations], [3.5, 1.4, 1.5])
    paragraph(doc, "Removing interval duration increased MAE by 155.02 kg, aircraft type by 29.18 kg, whole-flight timing by 16.77 kg, and time of day by 0.18 kg. TreeSHAP and permutation diagnostics also rank duration highest. These are predictive associations, not causal importance estimates, and only concern metadata.")
    figure(doc, "metadata_ablation.png", "Figure 3", "Increase in October MAE when each metadata feature group is removed and the boosted model is refit.", width=4.7)

    heading(doc, "6 Discussion")
    paragraph(doc, "The benchmark confirms that metadata alone supports a useful but incomplete estimator. The gap between boosting and linear baselines suggests nonlinear interactions in duration and aircraft type. The wide route-cluster interval warns against treating thousands of intervals as independent flight experiments. The near tie between a single boosted model and its tree ensemble shows that adding members is not automatically useful.")
    paragraph(doc, "PI-MCS targets a distinct operational problem: large-error estimates can be withheld under an explicit risk budget, while minimum release coverage and subgroup audits limit silent exclusion. That potential contribution is implemented as a decision system but is not yet demonstrated by the present tables. Only four October intervals belong to an aircraft type absent from training, so no unseen-aircraft generalization conclusion follows. Receiver selection, ACARS label noise, anonymized airlines, one year of observations, and inspected October labels further limit external validity.")
    paragraph(doc, "A completed confirmatory study requires full released trajectories aligned with weather and aircraft parameters, independently frozen labels, multiple years or genuinely independent datasets, and equal-budget comparisons of full PI-MCS against the selector and uncertainty baselines. The paper should report failures as well as wins. Nothing in the current evidence establishes real-time performance or Q1 journal acceptance.")

    heading(doc, "7 Conclusion")
    paragraph(doc, "This work specifies and implements a reproducible physics, uncertainty and risk-budget framework for aircraft fuel estimates. Verified PRC metadata experiments show strong nonlinear baselines and expose the dominant role of interval duration, but do not yet validate the proposed full decision system. The principal next scientific test is whether PI-MCS improves the error-versus-release trade-off under route and aircraft shift without reducing representation of difficult segments.")

    heading(doc, "Declarations")
    paragraph(doc, "Data availability. The PRC 2025 development data are available at https://doi.org/10.5281/zenodo.19184662 under the source record's access terms. Six input checksums and derived aggregate results are committed in the research extension. Additional OpenSky and ERA5 inputs require their own access and licence review.")
    paragraph(doc, "Code availability. The complete development implementation, fixed configuration, tests, reproducible result builder and manuscript source are committed on branch q1-research-extension of https://github.com/haarikaalla/flight-fuel-analysis. The upstream repository's redistribution status should be confirmed before a public submission archive is released.")
    paragraph(doc, "Ethics and consent. No human participants are studied; airline identities in the source data are anonymized. Author names, affiliations, funding, contributions, acknowledgements and competing-interest declarations are withheld from this anonymized review copy and require author confirmation before submission.")

    heading(doc, "References")
    references = [
        "[1] Sun J, Spinielli E, Strohmeier M (2026) Aircraft fuel burn estimation: the EUROCONTROL PRC 2025 data challenge. Journal of Open Aviation Science. https://doi.org/10.59490/joas.2026.8750",
        "[2] Grigoriou Y et al (2026) Physics-guided gradient boosting for aircraft fuel consumption estimation. Journal of Open Aviation Science. https://doi.org/10.59490/joas.2026.8764",
        "[3] Kačar A, Ćulibrk D (2026) Trajectory and meteorological modelling for aircraft fuel estimation. Journal of Open Aviation Science. https://doi.org/10.59490/joas.2026.8766",
        "[4] Romano Y, Patterson E, Candès EJ (2019) Conformalized quantile regression. NeurIPS. https://proceedings.neurips.cc/paper/2019/hash/5103c3584b063c431bd1268e9b5e76fb-Abstract.html",
        "[5] Tibshirani RJ et al (2019) Conformal prediction under covariate shift. NeurIPS. https://proceedings.neurips.cc/paper_files/paper/2019/hash/8fb21ee7a2207526da55a679f0332de2-Abstract.html",
        "[6] Gibbs I, Candès EJ (2021) Adaptive conformal inference under distribution shift. NeurIPS. https://proceedings.neurips.cc/paper/2021/hash/0d441de75945e5acbc865406fc9a2559-Abstract.html",
        "[7] Gibbs I, Cherian J, Candès EJ (2023) Conformal prediction with conditional guarantees. arXiv:2305.12616. https://arxiv.org/abs/2305.12616",
        "[8] Shah RD et al (2022) Selective regression under fairness constraints. ICML. https://proceedings.mlr.press/v162/shah22a.html",
        "[9] Bertsimas D, Sim M (2004) The price of robustness. Operations Research 52:35–53. https://doi.org/10.1287/opre.1030.0065",
        "[10] Angelopoulos AN et al (2024) Conformal risk control. ICLR. https://openreview.net/forum?id=33XGfHLtZg",
        "[11] Jarry G et al (2025) Domain generalization for aircraft fuel flow on unseen types. Transportation Research Part C. https://doi.org/10.1016/j.trc.2025.105143",
    ]
    for reference in references:
        doc.add_paragraph(reference)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("PI-MCS development manuscript  •  ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build())
